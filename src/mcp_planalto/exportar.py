"""Exportação de resultados para planilha (XLSX) ou CSV."""

from __future__ import annotations

import csv
import datetime as dt
import re
import secrets
from pathlib import Path

from .referencias import normalizar

COLUNAS_MAPA = [
    ("posicao", "#", 6),
    ("norma", "Norma", 34),
    ("data", "Data", 12),
    ("camada", "Camada", 26),
    ("relevancia", "Relevância", 11),
    ("ementa", "Ementa", 80),
    ("apelido", "Apelido", 30),
    ("situacao", "Situação", 30),
    ("motivos", "Por que entrou", 70),
    ("url", "Texto no Planalto", 60),
]


def _slug(s: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", normalizar(s)).strip("-")[:50] or "resultado"


def _valor(v):
    if isinstance(v, (list, tuple)):
        return " | ".join(str(x) for x in v)
    return v


def exportar(linhas: list[dict], destino_dir: Path, nome: str, formato: str = "xlsx",
             colunas: list[tuple[str, str, int]] | None = None, titulo: str | None = None,
             notas: list[str] | None = None) -> Path:
    colunas = colunas or COLUNAS_MAPA
    destino_dir.mkdir(parents=True, exist_ok=True)
    carimbo = dt.datetime.now().strftime("%Y%m%d-%H%M%S") + "-" + secrets.token_hex(2)
    caminho = destino_dir / f"{_slug(nome)}-{carimbo}.{formato}"
    if formato == "csv":
        with caminho.open("w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f, delimiter=";")
            w.writerow([c[1] for c in colunas])
            for l in linhas:
                w.writerow([_valor(l.get(c[0])) for c in colunas])
        return caminho
    if formato != "xlsx":
        raise ValueError("Formato deve ser 'xlsx' ou 'csv'.")

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Normas"
    linha0 = 1
    if titulo:
        ws.cell(1, 1, titulo).font = Font(bold=True, size=13)
        ws.cell(2, 1, f"Gerado em {dt.datetime.now():%d/%m/%Y %H:%M} por mcp-planalto "
                      "(fontes: Planalto e Senado Federal)").font = Font(italic=True, color="666666")
        linha0 = 4
    for j, (_k, rot, larg) in enumerate(colunas, 1):
        c = ws.cell(linha0, j, rot)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3864")
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(j)].width = larg
    for i, l in enumerate(linhas, linha0 + 1):
        for j, (k, _rot, _l) in enumerate(colunas, 1):
            v = _valor(l.get(k))
            c = ws.cell(i, j, v)
            c.alignment = Alignment(vertical="top", wrap_text=k in ("ementa", "motivos", "dispositivos"))
            if k == "url" and v:
                c.hyperlink = v
                c.font = Font(color="0563C1", underline="single")
    ws.freeze_panes = ws.cell(linha0 + 1, 3)
    ws.auto_filter.ref = f"A{linha0}:{get_column_letter(len(colunas))}{linha0 + len(linhas)}"
    if notas:
        ws2 = wb.create_sheet("Notas")
        for i, n in enumerate(notas, 1):
            ws2.cell(i, 1, n)
        ws2.column_dimensions["A"].width = 140
    wb.save(caminho)
    return caminho
