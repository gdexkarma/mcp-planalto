"""Testes do serviço sem rede: um cliente HTTP falso serve as fixtures."""

import pytest

from mcp_planalto.db import Relacao, consulta_fts
from mcp_planalto.exportar import exportar
from mcp_planalto.referencias import Referencia
from mcp_planalto.servico import (
    _consulta_regex,
    _ementa_altera,
    cita_norma,
    classificar_relacao,
    ler_data,
    normalizar_tipos,
    paginar,
)


def test_resolver_sem_ano(L):
    assert L.resolver("Lei 9.430").chave == "LEI:9430:1996"


def test_texto_e_paginacao(L):
    r = L.texto("Lei 9.430/96", dispositivo="art. 2, § 1")
    assert r["texto"] == "§ 1º A alíquota é de quinze por cento."
    assert r["planalto_atualizado_em"] == "2026-09-29"
    r = L.texto("Lei 9.430/96", max_caracteres=300, pagina=2)
    assert r["pagina"] == 2 and r["total_paginas"] > 2
    r = L.texto("Lei 9.430/96", termo="capital próprio")
    assert r["artigos_com_o_termo"] == 1
    r = L.texto("Lei 9.430/96", dispositivo="art. 99")
    assert "aviso" in r


def test_ficha_e_historico(L):
    f = L.ficha("Lei 9.430/1996")
    assert f["apelido"] == "Lei do Ajuste Tributário (1996)"
    assert f["ementa"] == "Dispõe sobre a legislação tributária federal."  # a do Planalto prevalece
    assert f["alterada_por"] == 2  # 12.973 e 15.525; a alteração vetada não conta
    h = L.historico("Lei 9.430/1996", dispositivo="art. 2")
    assert [a["norma"] for a in h["alteracoes"]] == ["Lei nº 15.525/2026", "Lei nº 12.973/2014"]
    h = L.historico("Lei 9.430/1996", dispositivo="art. 74, § 12")
    assert h["alteracoes"][0]["dispositivos"] == ["Alteração: Art. 74, § 12, Inciso 2, Alínea f"]
    h = L.historico("Lei 9.430/1996", direcao="feitas")
    assert [a["chave"] for a in h["alteracoes"]] == ["LEI:4502:1964"]  # correlata fica de fora


def test_verificar_atualizacao(L):
    r = L.verificar_atualizacao("Lei 9.430/1996")
    assert r["alteracoes_conhecidas"] == 2
    assert [p["norma"] for p in r["nao_refletidas_no_texto"]] == ["Lei nº 15.525/2026"]
    assert [p["norma"] for p in r["normas_recentes_que_citam_na_ementa"]] == ["Lei nº 15.600/2026"]
    assert r["conclusao"].startswith("ATENÇÃO")


def test_busca_sem_acento(L):
    r = L.buscar("legislacao tributaria")
    assert {x["chave"] for x in r["resultados"]} == {"LEI:9430:1996", "LEI:12973:2014"}
    r = L.buscar('"imposto sobre a renda das pessoas juridicas"')
    assert [x["chave"] for x in r["resultados"]] == ["LEI:12973:2014"]
    r = L.buscar("IRPJ OU compensação", tipos=["LEI"], ano_inicio=2014)
    assert {x["chave"] for x in r["resultados"]} == {"LEI:12973:2014", "LEI:15600:2026"}


def test_consulta_fts_segura():
    assert consulta_fts('lucro "real" OU IRPJ/CSLL') == '("lucro" AND "real") OR (("irpj" OR "csll"))'
    assert consulta_fts("transação tributária") == '("transac"* AND "tributari"*)'
    assert consulta_fts("tribut*") == '("tribut"*)'
    assert consulta_fts('"') == ""


def test_mapear_tema_livre(L):
    r = L.mapear_tema("tema de teste", termos_extras=['"lucro real"'], normas_extras=["Lei 9.430/1996"])
    por_chave = {n["chave"]: n for n in r["normas"]}
    assert por_chave["LEI:9430:1996"]["camada"] == "núcleo"
    assert por_chave["LEI:12973:2014"]["camada"] == "alteradora/regulamentadora"
    assert por_chave["LEI:15525:2026"]["camada"] == "alteradora/regulamentadora"
    assert "LEI:14000:2021" not in por_chave  # só vetos: não entra no acervo
    assert r["normas"][0]["chave"] == "LEI:9430:1996"


def test_exportar(tmp_path):
    linhas = [{"posicao": 1, "norma": "Lei nº 9.430/1996", "motivos": ["a", "b"], "url": "https://x"}]
    p = exportar(linhas, tmp_path, "Acervo IRPJ", "xlsx", titulo="Teste")
    from openpyxl import load_workbook

    ws = load_workbook(p).active
    assert ws["B5"].value == "Lei nº 9.430/1996" and ws["I5"].value == "a | b"
    c = exportar(linhas, tmp_path, "Acervo IRPJ", "csv")
    assert "Lei nº 9.430/1996" in c.read_text("utf-8-sig")


def test_paginar():
    p = paginar("\n".join(f"linha {i:04d} " + "x" * 40 for i in range(100)), 500, 3)
    assert p.pagina == 3 and p.total_paginas > 3 and p.texto.startswith("linha")


def test_classificar_relacao():
    R = lambda d, a="": Relacao("A", "B", d, a)  # noqa: E731
    assert classificar_relacao(R("Declaração de Alteração Permanente", "Alteração")) == "altera"
    assert classificar_relacao(R("Declaração de Alteração Permanente", "")) == "altera"  # sem itens
    assert classificar_relacao(R("Declaração de Revogação Permanente da Norma no Todo")) == "altera"
    assert classificar_relacao(R("Declaração de Alteração Permanente", "Alteração Vetada")) == "vetada"
    assert classificar_relacao(R("Declaração de Legislação Correlata", "Dispositivo Correlato")) == "correlata"
    assert classificar_relacao(R("Declaração de Alteração Permanente", "Ressalva")) == "ressalva"
    assert classificar_relacao(R("Declaração de Regulamentação", "Regulamentado")) == "regulamenta"


def test_cita_norma():
    lei = Referencia("LEI", "9430", 1996)
    assert cita_norma("(Redação dada pela Lei nº 9.430, de 1996)", lei)
    assert cita_norma("Lei n o 9.430, de 27 de dezembro de 1996", lei)
    assert cita_norma("altera as Leis nºs 9.249, de 1995, e 9.430, de 1996", lei)
    assert not cita_norma("Decreto-Lei nº 9.430", lei)
    assert not cita_norma("Lei Complementar nº 9.430", lei)
    assert cita_norma("Decreto-Lei nº 2.848, de 1940", Referencia("DEL", "2848", 1940))
    assert not cita_norma("Decreto-Lei nº 2.848, de 1940", Referencia("LEI", "2848", 1956))
    assert cita_norma("(Incluído pela Emenda Constitucional nº 45, de 2004)", Referencia("EMC", "45", 2004))
    assert cita_norma("(Redação dada pela Lei Complementar nº 227, de 2026)", Referencia("LCP", "227", 2026))
    assert cita_norma("(Incluído pela Medida Provisória nº 2.158-35, de 2001)", Referencia("MPV", "2158-35", 2001))


def test_ementa_altera():
    lei = Referencia("LEI", "12855", 2013)
    assert _ementa_altera("Altera a Lei nº 12.855, de 2 de setembro de 2013.", lei)
    assert not _ementa_altera("Altera o Decreto nº 8.000, que regulamenta a Lei nº 12.855, de 2013.", lei)


def test_consulta_regex_fronteira():
    pis = _consulta_regex("PIS")
    assert pis("contribuicao para o pis") and not pis("piso salarial") and not pis("piscicultura")
    assert _consulta_regex("tribut*")("tributacao")
    assert not _consulta_regex('"lucro real"')("lucro realizado")


def test_validacoes():
    assert ler_data("01/09/2026") == "2026-09-01" and ler_data("2026-09-01") == "2026-09-01"
    with pytest.raises(ValueError):
        ler_data("ontem")
    assert normalizar_tipos(["LC", "mp", "DL", "EC"]) == ["LCP", "MPV", "DEL", "EMC"]
    with pytest.raises(ValueError):
        normalizar_tipos(["XYZ"])


def test_historico_inclui_revogacao_da_norma_inteira(L):
    L.db.salvar_relacoes([Relacao("DEC:9580:2018", "LEI:9430:1996",
                                  "Declaração de Revogação Permanente da Norma no Todo", "", "", "2018-11-22")])
    L._detalhes.clear()
    det = L._detalhe_do_banco("LEI:9430:1996")
    assert any(r.origem == "DEC:9580:2018" for r in det.relacoes_recebidas)


def test_busca_vazia_nao_devolve_catalogo(L):
    for q in ["OR", "*", '"', "§"]:
        assert L.buscar(q)["total"] == 0
    r = L.buscar("Lei 9.430/96")
    assert r["resultados"][0]["chave"] == "LEI:9430:1996"


def test_mapear_nucleo_duplicado(L):
    r = L.mapear_tema("teste", normas_extras=["Lei 9.430/1996", "Lei 9.430/96"])
    assert [n["chave"] for n in r["normas"]].count("LEI:9430:1996") == 1
    assert r["normas"][0]["relevancia"] == 100


def test_mapear_escopo_por_dispositivo(L):
    r = L.mapear_tema("teste", normas_extras=["Lei 9.430/1996, art. 74"])
    chaves = {n["chave"] for n in r["normas"]}
    assert "LEI:12973:2014" in chaves  # alterou o art. 74
    assert "LEI:15525:2026" not in chaves  # só alterou o art. 2º
